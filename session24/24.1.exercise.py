# =============================================================
# EXERCITIUL 1  -  LISTA DE ANGAJATI IN EXCEL (scriere + citire)
# =============================================================
# Necesar:  pip install openpyxl
#
# CONTEXT (caz real)
# ------------------
# HR-ul vrea lista de angajati intr-un fisier .xlsx, ca sa o deschida in
# Excel. Tu o scrii din Python, apoi o citesti inapoi si calculezi
# totalul salariilor.
#
#
# CE AI DE FACUT
# --------------
# 1. scrie_angajati(cale, angajati):
#    - `angajati` e o lista de tupluri  (nume, salariu)
#    - creeaza un Workbook, redenumeste foaia activa "Angajati"
#    - scrie header-ul  ["nume", "salariu"]  apoi cate un rand per angajat
#    - salveaza la  cale
#
# 2. citeste_angajati(cale):
#    - deschide fisierul si intoarce o lista de tupluri  (nume, salariu),
#      SARIND peste randul de header (min_row=2)
#
#
# CERINTE
# -------
#   - foloseste  Workbook()  si  ws.append([...])
#   - salveaza cu  wb.save(cale)
#   - la citire foloseste  iter_rows(min_row=2, values_only=True)
#
#
# OUTPUT ASTEPTAT (din codul de test de mai jos)
# ----------------------------------------------
#   Ana 5000
#   Ion 7000
#   Maria 6500
#   total salarii: 18500
#
#
# INDICII
# -------
#   - scriere:  wb = Workbook(); ws = wb.active; ws.title = "Angajati"
#               ws.append(["nume", "salariu"]); ws.append([nume, salariu])
#   - citire:   for nume, salariu in ws.iter_rows(min_row=2, values_only=True): ...
# =============================================================

import os
from openpyxl import Workbook, load_workbook

TMP = "data_s24_ex1"
os.makedirs(TMP, exist_ok=True)


# ---- ZONA TA DE LUCRU ---------------------------------------

def scrie_angajati(cale, angajati):
    # TODO: creeaza workbook, scrie header + randuri, salveaza
    ...


def citeste_angajati(cale):
    # TODO: citeste (nume, salariu), fara header
    ...


# ---- COD DE TEST (nu trebuie sa-l modifici) -----------------
angajati = [("Ana", 5000), ("Ion", 7000), ("Maria", 6500)]

cale = os.path.join(TMP, "angajati.xlsx")
scrie_angajati(cale, angajati)

cititi = citeste_angajati(cale)
for nume, salariu in cititi:
    print(nume, salariu)

print("total salarii:", sum(s for _, s in cititi))    # 18500

