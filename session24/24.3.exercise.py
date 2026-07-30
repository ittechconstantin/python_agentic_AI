# =============================================================
# EXERCITIUL 3  (Excel + Streamlit)  -  JURNAL DE CHELTUIELI
# =============================================================
# Necesar:  pip install openpyxl streamlit
#
# CONTEXT (caz real)
# ------------------
# Vrei un jurnal personal de cheltuieli care RAMANE salvat intre
# rulari: adaugi cate o cheltuiala printr-un formular, ea se scrie
# intr-un fisier .xlsx de pe disc, iar mai jos vezi tot jurnalul de
# pana acum, plus totalul pe fiecare categorie.
#
# Combini doua lucruri pe care le stii deja:
#   - Excel (24.excel.py): Workbook(), load_workbook(), ws.append(...),
#     iter_rows(min_row=2, values_only=True), si agregarea pe categorie
#     din sectiunea 13 ("totaluri.get(categorie, 0) + suma").
#   - Streamlit (sesiunea 22): st.form + st.form_submit_button, ca sa
#     trimiti mai multe campuri deodata, la apasarea unui buton.
#
# DE CE E DIFERIT FATA DE UN SIMPLU SCRIPT PYTHON:
#   Streamlit reruleaza TOT scriptul de la inceput la fiecare
#   interactiune (fiecare click, fiecare submit de formular). Asta
#   inseamna ca fisierul .xlsx de pe disc e SINGURA ta "memorie" intre
#   doua rulari - de aceea la fiecare rerulare RECITESTI fisierul de pe
#   disc, ca sa arati mereu starea curenta (inclusiv ce tocmai ai
#   adaugat).
#
# DE CE NU FOLOSESTI MODUL "a" (append) CA LA CSV:
#   La fisiere text poti deschide cu open(cale, "a") si adaugi direct o
#   linie la final. La .xlsx nu exista asa ceva - un Workbook trebuie
#   mereu incarcat COMPLET in memorie, modificat, si SALVAT COMPLET la
#   loc. De aceea logica de salvare are doua ramuri:
#     - fisierul EXISTA deja pe disc -> il DESCHIZI cu load_workbook,
#       iei foaia lui activa, ii adaugi randul nou cu ws.append(...);
#     - fisierul NU exista inca (prima cheltuiala din tot jurnalul) ->
#       creezi un Workbook() NOU, scrii TU insuti header-ul
#       ["categorie", "suma"] pe primul rand (pentru ca acest Workbook
#       nu are inca niciun rand), apoi adaugi randul cu cheltuiala.
#   In AMBELE cazuri, la final apelezi wb.save(CALE_XLSX).
#
#
# CE AI DE FACUT  (scrie in ZONA TA DE LUCRU)
# -------------------------------------------
# 1. Afiseaza un titlu de pagina: "Jurnalul meu de cheltuieli".
#
# 2. Construieste un FORMULAR cu:
#      - un selectbox pentru categorie, cu optiunile din lista
#        CATEGORII (definita mai jos in cod);
#      - un camp numeric pentru suma (in lei), care nu trebuie sa
#        permita valori negative;
#      - un buton de trimitere.
#
# 3. Cand formularul a fost trimis SI suma completata e mai mare ca 0
#    (evita sa salvezi un rand gol daca userul apasa fara sa completeze
#    nimic):
#      - salveaza o cheltuiala noua (categorie + suma) in fisierul
#        CALE_XLSX;
#      - daca fisierul exista deja pe disc, adauga DOAR randul nou,
#        FARA sa rescrii header-ul; daca fisierul nu exista inca (prima
#        cheltuiala din tot jurnalul), creeaza-l si scrie mai intai
#        header-ul ["categorie", "suma"], apoi randul cu cheltuiala;
#      - anunta userul (ex: cu un mesaj scurt) ca s-a salvat.
#
# 4. INDIFERENT daca s-a trimis formularul sau nu (adica si la simpla
#    deschidere a paginii), afiseaza starea curenta a jurnalului:
#      - daca fisierul exista deja -> citeste toate cheltuielile
#        salvate (fara randul de header) si afiseaza-le ca tabel;
#      - calculeaza si afiseaza TOTALUL GENERAL cheltuit pana acum;
#      - calculeaza totalul PE CATEGORIE (o valoare per categorie) si
#        afiseaza fiecare categorie cu totalul ei;
#      - daca fisierul nu exista inca -> afiseaza un mesaj ca nu exista
#        nicio cheltuiala salvata.
#
#
# CERINTE
# -------
#   - from openpyxl import Workbook, load_workbook
#   - os.path.exists(CALE_XLSX) decide daca creezi fisier nou sau
#     deschizi cel existent - fara aceasta verificare, load_workbook ar
#     arunca eroare la prima rulare (cand fisierul inca nu exista)
#   - nu uita sa salvezi workbook-ul dupa fiecare adaugare - altfel
#     cheltuiala ramane doar in memorie si dispare la urmatoarea
#     rerulare a scriptului
#   - pasul 4 (afisarea) trebuie sa fie IN AFARA blocului conditionat de
#     trimiterea formularului de la pasul 3 - vrei sa vezi jurnalul
#     chiar si cand nu tocmai ai trimis formularul
#
#
# CUM VERIFICI
# ------------
#   streamlit run 24.3.exercise.py
#   adauga 2-3 cheltuieli din categorii diferite, apoi opreste si
#   reporneste scriptul - jurnalul trebuie sa fie tot acolo (e pe disc,
#   nu doar in memorie).
#
#
# COMPORTAMENT ASTEPTAT (exemplu)
# -------------------------------
#   adaugi Mancare 50, Transport 30, Mancare 80
#   -> tabelul arata cele 3 randuri
#   -> total general: 160 lei
#   -> Mancare: 130 lei, Transport: 30 lei
#
#
# INDICII
# -------
#   if trimis and suma > 0:
#       if os.path.exists(CALE_XLSX):
#           wb = load_workbook(CALE_XLSX)
#           ws = wb.active
#       else:
#           wb = Workbook()
#           ws = wb.active
#           ws.title = "Cheltuieli"
#           ws.append(["categorie", "suma"])
#       ws.append([categorie, suma])
#       wb.save(CALE_XLSX)
# =============================================================

import os
from openpyxl import Workbook, load_workbook
import streamlit as st

TMP = "data_s24_ex3"
os.makedirs(TMP, exist_ok=True)
CALE_XLSX = os.path.join(TMP, "cheltuieli.xlsx")

CATEGORII = ["Mancare", "Transport", "Utilitati", "Divertisment", "Altele"]


# ---- ZONA TA DE LUCRU ---------------------------------------
# TODO: scrie aici interfata (title + form + salvare in Excel +
#       afisare tabel + total general + total pe categorie)


