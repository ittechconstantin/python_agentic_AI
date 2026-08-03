# =============================================================
# EXERCITIUL 2  -  RAPORT CU DOUA FOI (multi-sheet)
# =============================================================
# Necesar:  pip install openpyxl
#
# CONTEXT (caz real)
# ------------------
# Un avantaj mare al Excel-ului fata de CSV: un singur fisier poate avea
# MAI MULTE FOI. Vrei un raport cu foaia "Useri" si foaia "Produse", apoi
# sa poti citi oricare foaie inapoi ca lista de dict-uri.
#
#
# CE AI DE FACUT
# --------------
# 1. exporta(cale, useri, produse):
#    - `useri`   = lista de tupluri  (nume, rol)
#    - `produse` = lista de tupluri  (nume, pret)
#    - creeaza un Workbook cu DOUA foi:
#        * foaia activa redenumita "Useri", cu header ["nume", "rol"]
#        * o foaie noua "Produse" (create_sheet), cu header ["nume", "pret"]
#      si scrie randurile corespunzatoare in fiecare
#    - salveaza la  cale
#
# 2. citeste_foaie(cale, nume_foaie):
#    - deschide fisierul, ia foaia  nume_foaie
#    - intoarce o lista de dict-uri, folosind PRIMUL rand ca header
#      (ex: {"nume": "Laptop", "pret": 4500})
#
#
# CERINTE
# -------
#   - foaia activa: ws = wb.active; ws.title = "Useri"
#   - a doua foaie:  wb.create_sheet("Produse")
#   - citirea ca dict-uri:  headers = next(randuri); dict(zip(headers, r))
#
#
# OUTPUT ASTEPTAT (din codul de test de mai jos)
# ----------------------------------------------
#   foi: ['Useri', 'Produse']
#   produse: [{'nume': 'Laptop', 'pret': 4500}, {'nume': 'Mouse', 'pret': 79}]
#
#
# INDICII
# -------
#   - dupa ws.title="Useri", scrii: ws.append(["nume","rol"]) + randurile
#   - ws2 = wb.create_sheet("Produse"); ws2.append(["nume","pret"]) + randuri
#   - citire:  randuri = ws.iter_rows(values_only=True); headers = next(randuri)
# =============================================================

import os
from openpyxl import Workbook, load_workbook
import os
from openpyxl import workbook, load_workbook


TMP = "data_s24_ex2"
os.makedirs(TMP, exist_ok=True)


# ---- ZONA TA DE LUCRU ---------------------------------------

def exporta(cale, useri, produse):
    # TODO: workbook cu foile "Useri" si "Produse", fiecare cu header + randuri
    wb = Workbook()
    ws = wb.active
    ws.title = "Useri"
    ws.append(["nume", "rol"])
    for nume, rol in useri:
        ws.append([nume, rol])

    ws2 = wb.create_sheet("Produse")
    ws2.append(["nume", "pret"])
    for nume, pret in produse:
        ws2.append([nume, pret])
    wb.save(cale)


def citeste_foaie(cale, nume_foaie):
    # TODO: intoarce lista de dict-uri din foaia ceruta (primul rand = header)
    ws = load_workbook(cale)[nume_foaie]
    randuri = ws.iter_rows(values_only=True)
    headers = next(randuri)
    return[dict(zip(headers, r))for r in randuri]



# ---- COD DE TEST (nu trebuie sa-l modifici) -----------------
useri = [("Ana", "admin"), ("Ion", "user")]
produse = [("Laptop", 4500), ("Mouse", 79)]

cale = os.path.join(TMP, "raport.xlsx")
exporta(cale, useri, produse)

print("foi:", load_workbook(cale).sheetnames)          # ['Useri', 'Produse']
print("produse:", citeste_foaie(cale, "Produse"))


